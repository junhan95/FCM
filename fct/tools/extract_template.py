"""
Frankonia Calculation Table — Excel 템플릿 추출기

엑셀(Global calculation table template_rev-x3.10.xlsx)을 읽어
  data/prices.json     : 가격 DB (Prices resume 시트, 계산된 값)
  data/templates.json  : 계산 시트(Total + 챔버/룸/옵션 시트)의 셀 모델(값/수식)
  data/expected.json   : 엑셀에 캐시된 계산 결과 (검증용)
을 생성한다.

사용법: python tools/extract_template.py <xlsx> <출력폴더>
"""
import sys, json, re, datetime
import openpyxl
from openpyxl.utils import get_column_letter

SRC = sys.argv[1]
OUT = sys.argv[2]

PRICES_SHEET = "Prices resume"

wb_f = openpyxl.load_workbook(SRC)                 # 수식
wb_v = openpyxl.load_workbook(SRC, data_only=True) # 캐시된 값


def fmt_class(nf: str) -> str:
    if nf is None or nf == "General":
        return "g"
    if "%" in nf:
        return "pct"
    if "€" in nf or "$€" in nf:
        return "eur"
    if "$" in nf:
        return "usd"
    if "0.000" in nf:
        return "d3"
    if "0.00" in nf:
        return "d2"
    if "0.0" in nf:
        return "d1"
    if nf.startswith("0") or "#,##0" in nf:
        return "int"
    if "yy" in nf.lower() or "mm" in nf.lower():
        return "date"
    return "g"


def norm_val(v):
    if isinstance(v, datetime.datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, datetime.date):
        return v.isoformat()
    return v


# ---------- 1. 가격 DB ----------
ws_v = wb_v[PRICES_SHEET]
ws_f = wb_f[PRICES_SHEET]
prices = []
category = None
cat_re = re.compile(r"^\s*(\d+)\s*-\s*(.+?)\s*$")
for r in range(1, ws_v.max_row + 1):
    rowvals = {}
    empty = True
    for c in range(1, 13):  # A..L
        v = norm_val(ws_v.cell(r, c).value)
        if v is not None:
            empty = False
        rowvals[get_column_letter(c)] = v
    if empty:
        continue
    cval = rowvals.get("C")
    if isinstance(cval, str) and rowvals.get("A") is None:
        m = cat_re.match(cval)
        if m:
            category = m.group(2)
    prices.append({
        "row": r,
        "ref": rowvals.get("A"),
        "sage": rowvals.get("B"),
        "item": rowvals.get("C"),
        "description": rowvals.get("D"),
        "unit": rowvals.get("E"),
        "price": rowvals.get("F"),
        "delivery": rowvals.get("G"),
        "colH": rowvals.get("H"),
        "colI": rowvals.get("I"),
        "colJ": rowvals.get("J"),
        "colK": rowvals.get("K"),
        "priceDate": rowvals.get("L"),
        "category": category if rowvals.get("A") is not None else None,
        "isHeader": rowvals.get("A") is None,
        "priceFormula": (ws_f.cell(r, 6).value if isinstance(ws_f.cell(r, 6).value, str) and str(ws_f.cell(r, 6).value).startswith("=") else None),
    })

# ---------- 2. 계산 시트 셀 모델 ----------
templates = {"sheets": {}, "order": [], "names": {}}
expected = {}

for name, dn in wb_f.defined_names.items():
    templates["names"][name] = dn.attr_text

for ws in wb_f.worksheets:
    if ws.title == PRICES_SHEET:
        continue
    wsv = wb_v[ws.title]
    cells = {}
    exp = {}
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if v is None:
                continue
            coord = c.coordinate
            entry = {}
            if isinstance(v, str) and v.startswith("="):
                entry["f"] = v[1:]
                cv = norm_val(wsv[coord].value)
                if cv is not None:
                    exp[coord] = cv
            else:
                entry["v"] = norm_val(v)
            fc = fmt_class(c.number_format)
            if fc != "g":
                entry["fmt"] = fc
            cells[coord] = entry
    templates["sheets"][ws.title] = {
        "title": ws.title,
        "hidden": ws.sheet_state != "visible",
        "maxRow": ws.max_row,
        "cells": cells,
        "merged": [str(m) for m in ws.merged_cells.ranges],
    }
    templates["order"].append(ws.title)
    expected[ws.title] = exp

# 데이터 유효성(드롭다운) — Total
dv = []
for d in wb_f["Total"].data_validations.dataValidation:
    dv.append({"sqref": str(d.sqref), "type": d.type, "formula1": d.formula1})
templates["sheets"]["Total"]["validations"] = dv

import os
os.makedirs(OUT, exist_ok=True)
with open(os.path.join(OUT, "prices.json"), "w", encoding="utf-8") as f:
    json.dump(prices, f, ensure_ascii=False, indent=0)
with open(os.path.join(OUT, "templates.json"), "w", encoding="utf-8") as f:
    json.dump(templates, f, ensure_ascii=False)
with open(os.path.join(OUT, "expected.json"), "w", encoding="utf-8") as f:
    json.dump(expected, f, ensure_ascii=False)

print("prices:", len(prices), "sheets:", len(templates["order"]),
      "formula cells:", sum(len(v) for v in expected.values()))
